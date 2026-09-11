from datetime import date
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from src.models import base
from src.services.db_service import DatabaseService
from src.services.export_service import ExportService
from src.services.import_service import ImportService
from src.services.machine_service import MachineService
from src.services.project_service import ProjectService
from src.services.user_service import UserService
from src.services.expense_service import ExpenseService
from src.ui.widgets.common import PaginationBar


def test_packaged_v1_uses_versioned_empty_production_directory(tmp_path, monkeypatch):
    import src.config as config

    monkeypatch.delenv("DEPLOYMATE_DATA_DIR", raising=False)
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    monkeypatch.setattr(config.sys, "frozen", True, raising=False)
    monkeypatch.setattr(config.sys, "platform", "win32")

    database_path = config.resolve_db_path()

    assert config.get_db_mode() == "prod"
    assert Path(database_path).parent.parent == tmp_path / "DeployMate"
    assert Path(database_path).parent.name.startswith("V1.1-")
    assert not (tmp_path / "DeployMate" / "deploymate.db").exists()


def test_requested_exports_keep_full_machine_data_and_zero_values(tmp_path):
    DatabaseService(str(tmp_path / "requested-export.db"))
    project = ProjectService().create_project(
        "完整导出项目", customer="客户乙", remarks="说明行内容"
    )
    machine = MachineService().create_machine(
        project.id,
        "10.40.0.10",
        role="master",
        business_ip="10.40.0.11",
        gpu_count="0",
        remarks="机器备注",
        extra_info='{"SSH端口": 2222, "厂商": "厂商乙"}',
    )
    service = ExportService()
    word_path = tmp_path / "full.docx"
    markdown_path = tmp_path / "full.md"
    excel_path = tmp_path / "full.xlsx"
    service.export_projects([project.id], str(word_path), "docx")
    service.export_projects([project.id], str(markdown_path), "md")
    service.export_projects([project.id], str(excel_path), "xlsx")

    document = Document(word_path)
    word_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    word_text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    markdown = markdown_path.read_text(encoding="utf-8")
    workbook = load_workbook(excel_path, data_only=True)
    machine_headers = [cell.value for cell in workbook["机器"][1]]
    machine_values = [cell.value for cell in workbook["机器"][2]]

    assert "10.40.0.10" in word_text
    assert "0" in word_text
    assert "机器备注" in markdown
    assert "SSH端口" in markdown
    assert "备注" in machine_headers
    assert "其他信息" in machine_headers
    assert machine_values[machine_headers.index("GPU数量")] == "0"


def test_requested_import_preserves_csv_rows_and_pagination_boundary(tmp_path):
    DatabaseService(str(tmp_path / "requested-import.db"))
    csv_path = tmp_path / "machines.csv"
    csv_path.write_text(
        "管理网IP,角色,未知字段\n10.50.0.10,管理节点,原始值\n",
        encoding="utf-8",
    )
    groups = ImportService().read_file_groups(str(csv_path))
    assert groups[0]["source_rows"][0]["未知字段"] == "原始值"
    assert groups[0]["rows"][0]["extra_info"] == '{"未知字段": "原始值"}'

    pager = PaginationBar(page_size=20)
    pager.set_total(41)
    pager.set_page(3)
    assert pager.page == 3
    assert pager.page_count == 3
    pager.set_total(21)
    assert pager.page == 2
    assert pager.page_count == 2
    assert "第 2 / 2 页" in pager.summary.text()


def test_local_user_login_password_change_and_disable(tmp_path):
    DatabaseService(str(tmp_path / "users.db"))
    service = UserService()
    assert service.has_users() is False
    admin = service.create_user("admin", "secret123", "管理员", role="admin")
    assert service.authenticate("admin", "wrong") is None
    assert service.authenticate("admin", "secret123").username == "admin"
    service.change_password(admin.id, "secret123", "newsecret123")
    assert service.authenticate("admin", "secret123") is None
    assert service.authenticate("admin", "newsecret123").username == "admin"
    guest = service.create_user("guest", "guest123", "访客")
    service.set_active(guest.id, False)
    assert service.authenticate("guest", "guest123") is None
    service.set_active(guest.id, True)
    assert service.authenticate("guest", "guest123").username == "guest"
    service.set_active(guest.id, False)
    try:
        service.set_active(admin.id, False)
        raise AssertionError("at least one active user should remain")
    except ValueError as exc:
        assert "至少保留一个启用用户" in str(exc)
    assert service.is_admin(service.authenticate("admin", "newsecret123")) is True


def test_expense_profit_uses_assessment_fee_minus_actual_travel_cost(tmp_path):
    DatabaseService(str(tmp_path / "profit.db"))
    project = ProjectService().create_project("利润项目")
    expenses = ExpenseService()
    expenses.create_expense(
        project.id, date(2026, 9, 10), "出差考核费用", "1000",
        travel_person="张三", travel_start_date=date(2026, 9, 8),
        travel_end_date=date(2026, 9, 10),
    )
    expenses.create_expense(
        project.id, date(2026, 9, 10), "交通", "120",
        travel_person="张三", travel_start_date=date(2026, 9, 8),
        travel_end_date=date(2026, 9, 10),
    )
    expenses.create_expense(
        project.id, date(2026, 9, 10), "住宿", "80",
        travel_person="张三", travel_start_date=date(2026, 9, 8),
        travel_end_date=date(2026, 9, 10),
    )
    summary = expenses.get_profit_summary(project.id)
    assert summary["assessment_fee"] == 1000.0
    assert summary["actual_cost"] == 200.0
    assert summary["profit"] == 800.0
    assert expenses.get_total_amount(project.id) == 200.0

    saved = expenses.get_all_expenses()
    assert saved[0].travel_person == "张三"
    assert saved[0].travel_start_date == date(2026, 9, 8)
    assert saved[0].travel_end_date == date(2026, 9, 10)


def test_expense_category_separates_assessment_and_travel_costs(tmp_path):
    DatabaseService(str(tmp_path / "expense-category.db"))
    project = ProjectService().create_project("费用分类项目")
    expenses = ExpenseService()
    assessment = expenses.create_expense(
        project.id, date(2026, 9, 11), "考核费用", "1000",
        expense_category="考核费用",
    )
    travel = expenses.create_expense(
        project.id, date(2026, 9, 11), "交通", "260",
        expense_category="出差费用",
    )

    assert assessment.expense_category == "考核费用"
    assert assessment.expense_type == "考核费用"
    assert travel.expense_category == "出差费用"
    assert expenses.get_profit_summary(project.id) == {
        "assessment_fee": 1000.0,
        "actual_cost": 260.0,
        "profit": 740.0,
        "reimbursed": 0.0,
        "count": 2,
    }
    output = tmp_path / "expense-category.xlsx"
    expenses.export_expenses(project.id, str(output))
    sheet = load_workbook(output, data_only=True).active
    assert "费用分类" in [cell.value for cell in sheet[1]]


def test_expense_summary_returns_one_total_row_per_project(tmp_path):
    DatabaseService(str(tmp_path / "expense-summary-projects.db"))
    first = ProjectService().create_project("项目甲")
    second = ProjectService().create_project("项目乙")
    expenses = ExpenseService()
    expenses.create_expense(first.id, date(2026, 9, 11), "考核费用", "2000", expense_category="考核费用")
    expenses.create_expense(first.id, date(2026, 9, 11), "交通", "1000", expense_category="出差费用")
    expenses.create_expense(second.id, date(2026, 9, 11), "考核费用", "500", expense_category="考核费用")
    expenses.create_expense(second.id, date(2026, 9, 11), "住宿", "120", expense_category="出差费用")

    summaries = {
        item["project_name"]: item
        for item in expenses.get_profit_summaries_by_project()
    }

    assert summaries["项目甲"]["assessment_fee"] == 2000.0
    assert summaries["项目甲"]["actual_cost"] == 1000.0
    assert summaries["项目甲"]["profit"] == 1000.0
    assert summaries["项目甲"]["count"] == 2
    assert summaries["项目乙"]["assessment_fee"] == 500.0
    assert summaries["项目乙"]["actual_cost"] == 120.0
    assert summaries["项目乙"]["profit"] == 380.0


def test_expense_profit_can_be_negative_when_travel_cost_exceeds_assessment(tmp_path):
    DatabaseService(str(tmp_path / "negative-profit.db"))
    project = ProjectService().create_project("负利润项目")
    expenses = ExpenseService()
    expenses.create_expense(
        project.id, date(2026, 9, 11), "考核费用", "100",
        expense_category="考核费用",
    )
    expenses.create_expense(
        project.id, date(2026, 9, 11), "交通", "260",
        expense_category="出差费用",
    )

    summary = expenses.get_profit_summary(project.id)

    assert summary["assessment_fee"] == 100.0
    assert summary["actual_cost"] == 260.0
    assert summary["profit"] == -160.0


def test_expense_travel_dates_are_validated_and_exported(tmp_path):
    DatabaseService(str(tmp_path / "expense-fields.db"))
    project = ProjectService().create_project("出差字段项目")
    expenses = ExpenseService()
    try:
        expenses.create_expense(
            project.id, date(2026, 9, 10), "交通", "20",
            travel_start_date=date(2026, 9, 12),
            travel_end_date=date(2026, 9, 11),
        )
        raise AssertionError("invalid travel dates should fail")
    except ValueError as exc:
        assert "出差结束不能早于出差开始" in str(exc)

    expense = expenses.create_expense(
        project.id, date(2026, 9, 10), "交通", "20",
        travel_person="李四",
        travel_start_date=date(2026, 9, 10),
        travel_end_date=date(2026, 9, 12),
    )
    output = tmp_path / "expenses.xlsx"
    expenses.export_expenses(project.id, str(output))
    workbook = load_workbook(output, data_only=True)
    headers = [cell.value for cell in workbook["出差费用"][1]]
    values = [cell.value for cell in workbook["出差费用"][2]]
    assert headers[:4] == ["费用日期", "出差人员", "出差开始", "出差结束"]
    assert values[1] == "李四"
    assert values[2].strftime("%Y-%m-%d") == "2026-09-10"

    word_output = tmp_path / "expenses.docx"
    markdown_output = tmp_path / "expenses.md"
    export_service = ExportService()
    export_service.export_projects([project.id], str(word_output), "docx")
    export_service.export_projects([project.id], str(markdown_output), "md")
    word_text = "\n".join(
        cell.text
        for table in Document(word_output).tables
        for row in table.rows
        for cell in row.cells
    )
    markdown_text = markdown_output.read_text(encoding="utf-8")
    assert all(field in word_text for field in ["出差人员", "出差开始", "出差结束", "李四"])
    assert all(field in markdown_text for field in ["出差人员", "出差开始", "出差结束", "李四"])
