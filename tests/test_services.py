from datetime import date
from pathlib import Path
import tempfile
import sys

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook, load_workbook
from docx import Document
from zipfile import ZipFile

from src.services.db_service import DatabaseService
from src.services.project_service import ProjectService
from src.services.machine_service import MachineService
from src.services.daily_report_service import DailyReportService
from src.services.expense_service import ExpenseService
from src.services.import_service import ImportService
from src.config import resolve_db_path
from src.services.ssh_service import SSHService
from src.services.export_service import ExportService
from src.utils.field_mapper import normalize_machine_header


def run():
    assert normalize_machine_header("节点\n类型") == ["role"]
    assert normalize_machine_header("业务网络地址") == ["business_ip"]
    assert normalize_machine_header("登录账号") == ["username"]
    assert normalize_machine_header("登录密码") == ["password"]
    assert normalize_machine_header("显卡数量") == ["gpu_count"]

    temp_db = Path(tempfile.gettempdir()) / "deploymate_test.db"
    if temp_db.exists():
        temp_db.unlink()

    db = DatabaseService(str(temp_db))
    session = db.get_session()
    session.close()

    workbook = Workbook()
    first = workbook.active
    first.title = "项目甲"
    first.append(["角色", "业务网IP", "用户名", "GPU数量"])
    first.append(["计算节点", "10.10.1.10", "root", 2])
    second = workbook.create_sheet("项目乙")
    second.append(["角色", "集群网IP", "用户名", "GPU型号"])
    second.append(["存储节点", "10.10.2.10", "deploy", "A800"])
    multi_sheet_file = Path(tempfile.gettempdir()) / "deploymate_multi_sheet.xlsx"
    workbook.save(multi_sheet_file)
    sheet_import_service = ImportService()
    groups = sheet_import_service.read_file_groups(multi_sheet_file)
    assert [group["sheet_name"] for group in groups] == ["项目甲", "项目乙"]
    assert len(groups[0]["rows"]) == len(groups[1]["rows"]) == 1
    for index, group in enumerate(groups):
        project_id = sheet_import_service.get_or_create_import_project(
            group["sheet_name"], multi_sheet_file.stem, index
        )
        imported = sheet_import_service.import_machines(
            project_id,
            group["rows"],
            multi_sheet_file.name,
            "xlsx",
            multi_sheet_file.stat().st_size,
        )
        created_project = ProjectService().get_project_by_id(project_id)
        assert created_project.name == group["sheet_name"]
        assert imported["new"] == 1
    multi_sheet_file.unlink(missing_ok=True)

    realistic_book = Workbook()
    realistic_sheet = realistic_book.active
    realistic_sheet.title = "SCM"
    realistic_sheet.merge_cells("A1:G1")
    realistic_sheet["A1"] = "SCM"
    realistic_sheet.append(["角色", "业务网IP", "用户名", "密码", "GPU数量", "ssh端口号", "备注"])
    realistic_sheet.append(["管理", "10.103.55.59", "root", "admin@2024", 8, 222, ""])
    realistic_sheet.append(["计算节点", "10.103.55.62", "root", "admin@2025", 8, 220, "2025年11月5日新加入"])
    notes_sheet = realistic_book.create_sheet("实施说明")
    notes_sheet.append(["实施时间", "销售"])
    notes_sheet.append(["2025/12/11", "徐明生"])
    empty_machine_sheet = realistic_book.create_sheet("待补录")
    empty_machine_sheet.append(["角色", "业务网IP", "用户名", "GPU数量"])
    empty_machine_sheet.append(["计算节点", "", "root", 8])
    realistic_file = Path(tempfile.gettempdir()) / "deploymate_realistic_multi_sheet.xlsx"
    realistic_book.save(realistic_file)
    realistic_groups = sheet_import_service.read_file_groups(realistic_file)
    realistic_sheet_names = {group["sheet_name"] for group in realistic_groups}
    assert {"SCM", "实施说明"}.issubset(realistic_sheet_names)
    assert len(realistic_groups[0]["rows"]) == 2
    assert realistic_groups[0]["project_metadata"]["sales"] == "徐明生"
    imported_project_id = sheet_import_service.get_or_create_import_project(
        "SCM", realistic_file.stem, 0, realistic_groups[0]["project_metadata"]
    )
    imported_project = ProjectService().get_project_by_id(imported_project_id)
    assert imported_project.project_code is None
    assert imported_project.sales == "徐明生"
    assert realistic_groups[0]["rows"][0]["ip"] == "10.103.55.59"
    assert "2025年11月5日新加入" in realistic_groups[0]["rows"][1]["remarks"]
    realistic_file.unlink(missing_ok=True)

    project_service = ProjectService()
    machine_service = MachineService()
    daily_service = DailyReportService()
    expense_service = ExpenseService()
    import_service = ImportService()

    project = project_service.create_project(
        name="测试项目",
        customer="测试客户",
        project_code="T-001",
        location="上海",
        start_date=date(2026, 9, 8),
        end_date=date(2026, 9, 9),
        sales="销售甲",
        status="实施中",
    )

    created = project_service.create_project(
        name="测试项目2",
        customer="测试客户2",
        project_code="T-002",
        location="北京",
        start_date=date(2026, 9, 10),
        end_date=date(2026, 9, 11),
        sales="销售乙",
        status="已完成",
    )
    assert created.project_code == "T-002"
    assert created.sales == "销售乙"
    updated = project_service.update_project(
        created.id,
        name="测试项目2-更新",
        customer="测试客户2",
        project_code="T-002",
        location="北京",
        start_date=date(2026, 9, 10),
        end_date=date(2026, 9, 12),
        sales="销售丙",
        status="实施中",
    )
    assert updated.name == "测试项目2-更新"
    assert updated.end_date == date(2026, 9, 12)
    try:
        project_service.update_project(
            created.id,
            start_date=date(2026, 9, 20),
            end_date=date(2026, 9, 19),
        )
        raise AssertionError("invalid updated project date range should fail")
    except ValueError as exc:
        assert "不能早于" in str(exc)
    try:
        project_service.create_project(
            name="X" * 101,
            customer="测试客户4",
            project_code="T-003",
            location="杭州",
            start_date=date(2026, 9, 14),
            end_date=date(2026, 9, 15),
            sales="销售丁",
            status="待开始",
        )
        raise AssertionError("oversized project name should fail")
    except ValueError as exc:
        assert "不能超过" in str(exc)
    try:
        project_service.create_project(
            name="测试项目3",
            customer="测试客户3",
            project_code="T-001",
            location="深圳",
            start_date=date(2026, 9, 12),
            end_date=date(2026, 9, 13),
            sales="销售戊",
            status="待开始",
        )
        raise AssertionError("duplicate project_code should fail")
    except ValueError as exc:
        assert "不能重复添加" in str(exc)

    machine = machine_service.create_machine(
        project_id=project.id,
        ip="10.0.0.1",
        role="GPU节点",
        business_ip="172.16.0.1",
        cluster_ip="10.20.0.1",
        compute_ip="10.30.0.1",
        storage_ip="10.40.0.1",
        username="root",
        account="root",
        password="123456",
        gpu_count="2",
        gpu_model="A800",
        hostname="node-01",
        os="Ubuntu 22.04",
        cpu="64C",
        memory="256G",
        gpu="A800 x 2",
        cuda="12.4",
        docker="27.3",
    )
    assert machine.project_id == project.id
    try:
        machine_service.create_machine(
            project_id=project.id,
            ip="10.0.0.99",
            role="缺少网络的节点",
        )
        raise AssertionError("machine without business/cluster network should fail")
    except ValueError as exc:
        assert "至少填写一个" in str(exc)

    merged, is_new, conflicts = machine_service.merge_machine_data(
        project_id=project.id,
        ip="10.0.0.1",
        new_data={"gpu_model": "A800", "gpu_count": "4"},
    )

    assert merged.id == machine.id
    assert is_new is False
    assert isinstance(conflicts, dict)
    assert machine_service.get_machine_by_id(machine.id).ip == "10.0.0.1"

    report = daily_service.create_report(
        project_id=project.id,
        report_date=date(2026, 9, 8),
        work_content="完成部署。",
        problems="无",
        solutions="无",
        next_plan="继续联调",
        status="已完成",
    )
    assert report.project_id == project.id

    report2 = daily_service.create_report(
        project_id=project.id,
        report_date=date(2026, 9, 8),
        work_content="完成联调。",
        problems="无",
        solutions="无",
        next_plan="收尾",
        status="已完成",
    )
    assert report.id == report2.id
    assert report2.work_content == "完成联调。"
    report_next = daily_service.create_report(
        project_id=project.id,
        report_date=date(2026, 9, 9),
        work_content="完成验收准备。",
        status="草稿",
    )
    assert daily_service.get_report_by_id(report_next.id).status == "草稿"
    try:
        daily_service.update_report(report_next.id, report_date=date(2026, 9, 8))
        raise AssertionError("duplicate report date should fail")
    except ValueError as exc:
        assert "已存在SOP记录" in str(exc)

    expense = expense_service.create_expense(
        project_id=project.id,
        expense_date=date(2026, 9, 8),
        expense_type="交通",
        amount="100",
        remarks="打车",
        is_reimbursed="未提交",
    )
    assert float(expense.amount) == 100.0
    assert expense_service.get_total_amount(project.id) == 100.0
    updated_expense = expense_service.update_expense(expense.id, amount="128.50", is_reimbursed="已提交")
    assert float(updated_expense.amount) == 128.5
    assert expense_service.get_expense_by_id(expense.id).is_reimbursed == "已提交"

    xlsx_path = Path(tempfile.gettempdir()) / "deploymate_test.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["IP地址", "角色", "业务网IP", "GPU型号"])
    ws.append(["10.0.0.2", "计算节点", "172.16.0.2", "A800"])
    wb.save(xlsx_path)

    parsed = import_service.read_file(str(xlsx_path))
    assert parsed and parsed[0]["ip"] == "10.0.0.2"
    assert parsed[0]["role"] == "计算节点"
    assert parsed[0]["business_ip"] == "172.16.0.2"
    imported = import_service.import_machines(
        project.id, parsed, xlsx_path.name, "xlsx", xlsx_path.stat().st_size
    )
    assert imported["new"] == 1
    assert imported["project_id"] == project.id
    assert imported["project_name"] == project.name
    assert machine_service.get_machine_by_ip(project.id, "10.0.0.2") is not None
    try:
        imported_machine = machine_service.get_machine_by_ip(project.id, "10.0.0.2")
        machine_service.update_machine(imported_machine.id, ip="10.0.0.1")
        raise AssertionError("duplicate machine IP update should fail")
    except ValueError as exc:
        assert "相同管理网IP" in str(exc)
    imported_again = import_service.import_machines(
        project.id, parsed, xlsx_path.name, "xlsx", xlsx_path.stat().st_size
    )
    assert imported_again["merge"] == 1
    history = import_service.get_import_history()
    assert history[0]["project"].startswith("测试项目")
    assert history[0]["file_name"] == xlsx_path.name
    grouped_xlsx = Path(tempfile.gettempdir()) / "deploymate_grouped_test.xlsx"
    grouped_wb = Workbook()
    grouped_ws = grouped_wb.active
    grouped_ws.append(["长水务AI开放平台"])
    grouped_ws.append(["角色", "业务网IP", "存储+集群网", "计算网", "用户名", "密码", "GPU数量", "GPU型号"])
    grouped_ws.append(["训练节点", "万兆网口1：10.6.90.3", "万兆网口2：192.168.33.11", "400G网络：192.168.40.11", "root", "secret", 8, "910B"])
    grouped_ws.append(["长水务AI开放平台-第二组"])
    grouped_ws.append(["角色", "业务网IP", "存储+集群网", "计算网", "用户名", "密码", "GPU数量", "GPU型号"])
    grouped_ws.append(["推理节点", "万兆网口1：10.6.90.6", "万兆网口2：192.168.33.21", "400G网络：192.168.40.21", "root", "secret", 2, "无"])
    grouped_wb.save(grouped_xlsx)
    grouped_rows = import_service.read_file(str(grouped_xlsx))
    assert len(grouped_rows) == 2
    assert grouped_rows[0]["ip"] == "10.6.90.3"
    assert grouped_rows[0]["cluster_ip"] == "192.168.33.11"
    assert grouped_rows[0]["compute_ip"] == "192.168.40.11"
    assert "gpu_model" not in grouped_rows[1]
    grouped_result = import_service.import_machines(
        project.id, grouped_rows, grouped_xlsx.name, "xlsx", grouped_xlsx.stat().st_size
    )
    assert grouped_result["new"] == 2
    assert grouped_result["errors"] == []
    assert machine_service.get_machine_by_ip(project.id, "10.6.90.3").gpu_model == "910B"
    assert SSHService.detect_gpu_interconnect("GPU0 GPU1\nGPU0 NVLink GPU1") == "NVLink"
    assert SSHService.detect_gpu_interconnect("GPU0 GPU1\nGPU0 PIX GPU1") == "PCIe"
    interfaces = SSHService.parse_network_interfaces(
        "2: eth0    inet 10.0.0.10/24 brd 10.0.0.255 scope global eth0\n"
        "3: ens3@if2    inet 192.168.1.10/24 brd 192.168.1.255 scope global ens3\n"
        "4: bond0    inet 172.16.0.10/16 brd 172.16.255.255 scope global bond0"
    )
    assert interfaces == [
        {"name": "eth0", "ip": "10.0.0.10", "prefix": 24},
        {"name": "ens3", "ip": "192.168.1.10", "prefix": 24},
        {"name": "bond0", "ip": "172.16.0.10", "prefix": 16},
    ]

    export_service = ExportService()
    export_dir = Path(tempfile.gettempdir()) / "deploymate_export_test"
    export_dir.mkdir(exist_ok=True)
    xlsx_export = export_dir / "project.xlsx"
    docx_export = export_dir / "project.docx"
    zip_export = export_dir / "project.zip"
    expense_export = export_dir / "expenses.xlsx"
    export_service.export_project(project.id, str(xlsx_export), "xlsx")
    export_service.export_project(project.id, str(docx_export), "docx")
    export_service.export_project(project.id, str(zip_export), "zip")
    expense_service.export_expenses(project.id, str(expense_export))
    assert load_workbook(xlsx_export, read_only=True).sheetnames == ["项目", "机器", "SOP问题", "出差费用"]
    assert Document(docx_export).paragraphs
    expense_book = load_workbook(expense_export, read_only=True)
    assert expense_book.sheetnames == ["出差费用"]
    assert expense_book["出差费用"].max_row == 2
    with ZipFile(zip_export) as archive:
        assert set(archive.namelist()) == {
            "01_项目数据/项目数据.xlsx",
            "02_实施资料/项目实施资料.docx",
            "03_SOP问题记录/SOP问题记录.md",
        }

    assert resolve_db_path(None, "dev").endswith("deploymate_dev.db")
    assert resolve_db_path(None, "prod").endswith("deploymate.db")

    disposable = project_service.create_project(
        name="待删除项目",
        customer="测试客户",
        project_code="DELETE-001",
    )
    machine_service.create_machine(
        project_id=disposable.id,
        ip="10.99.0.1",
        business_ip="172.99.0.1",
    )
    assert project_service.delete_project(disposable.id) is True
    assert project_service.get_project_by_id(disposable.id) is None
    assert machine_service.get_machine_by_ip(disposable.id, "10.99.0.1") is None

    print("OK")


if __name__ == "__main__":
    run()
