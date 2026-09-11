from typing import List
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import cast, func, or_, String
from src.models import Expense, Project, get_session
from src.services.access_service import owner_id_for, scope_query
from src.services.user_service import UserService
from src.utils.validators import (
    EXPENSE_FIELD_LIMITS,
    ValidationError,
    optional_text,
    require_text,
    validate_choice,
    validate_positive_amount,
    validate_date_range,
)

class ExpenseService:
    EXPENSE_CATEGORIES = {"出差费用", "考核费用"}
    EXPENSE_TYPES = {"交通", "住宿", "餐费", "打车", "其他", "费用", "出差费用", "考核费用", "出差考核费用"}
    ASSESSMENT_TYPES = {"考核费用", "出差考核费用"}

    def __init__(self, current_user=None):
        self.current_user = current_user

    def _query(self, session):
        return scope_query(session.query(Expense), Expense, self.current_user)

    def _project(self, session, project_id):
        return scope_query(session.query(Project), Project, self.current_user).filter(
            Project.id == int(project_id)
        ).first()

    @classmethod
    def _normalise_category_type(cls, expense_category, expense_type):
        expense_type = validate_choice(expense_type, "费用类型", cls.EXPENSE_TYPES)
        if expense_category is None or str(expense_category).strip() == "":
            expense_category = "考核费用" if expense_type in cls.ASSESSMENT_TYPES else "出差费用"
        else:
            expense_category = validate_choice(
                expense_category, "费用分类", cls.EXPENSE_CATEGORIES
            )
        if expense_category == "考核费用" and expense_type not in cls.ASSESSMENT_TYPES:
            expense_type = "考核费用"
        elif expense_category == "出差费用" and expense_type in cls.ASSESSMENT_TYPES:
            expense_type = "出差费用"
        return expense_category, expense_type

    @classmethod
    def _is_assessment(cls, expense):
        return (
            expense.expense_category == "考核费用"
            or expense.expense_type in cls.ASSESSMENT_TYPES
        )

    @classmethod
    def amount_breakdown(cls, expense) -> dict:
        if expense.assessment_fee is not None or expense.actual_cost is not None:
            assessment_fee = float(expense.assessment_fee or 0)
            actual_cost = float(expense.actual_cost or 0)
            return {
                "assessment_fee": assessment_fee,
                "actual_cost": actual_cost,
                "profit": assessment_fee - actual_cost,
            }
        amount = float(expense.amount or 0)
        assessment_fee = amount if cls._is_assessment(expense) else 0.0
        actual_cost = 0.0 if cls._is_assessment(expense) else amount
        return {
            "assessment_fee": assessment_fee,
            "actual_cost": actual_cost,
            "profit": assessment_fee - actual_cost,
        }

    def _filtered_expense_query(self, session, project_id=None, search="", owner_id=None):
        query = self._query(session)
        if project_id:
            query = query.filter(Expense.project_id == project_id)
        if owner_id and UserService.is_admin(self.current_user):
            query = query.filter(Expense.owner_id == int(owner_id))
        keyword = (search or "").strip()
        if keyword:
            pattern = f"%{keyword}%"
            query = query.outerjoin(Project, Project.id == Expense.project_id).filter(or_(
                Project.name.ilike(pattern),
                *[
                    cast(getattr(Expense, field), String).ilike(pattern)
                    for field in (
                        "expense_date", "travel_person", "travel_start_date",
                        "travel_end_date", "expense_category", "expense_type",
                        "amount", "remarks", "is_reimbursed",
                    )
                ],
            ))
        return query

    def get_expenses_page(self, page=1, page_size=20, project_id=None, search="", owner_id=None):
        session = get_session()
        try:
            query = self._filtered_expense_query(session, project_id, search, owner_id)
            query = query.order_by(Expense.expense_date.desc())
            total = query.with_entities(Expense.id).order_by(None).count()
            rows = query.outerjoin(Project, Project.id == Expense.project_id).with_entities(
                Expense, Project.name.label("project_name")
            ).offset((page - 1) * page_size).limit(page_size).all()
            return total, rows
        finally:
            session.close()
    def get_expenses_by_project(self, project_id: int) -> List[Expense]:
        session = get_session()
        try:
            return (
                self._query(session)
                .filter(Expense.project_id == project_id)
                .order_by(Expense.expense_date.desc())
                .all()
            )
        finally:
            session.close()

    def get_all_expenses(self) -> List[Expense]:
        session = get_session()
        try:
            return self._query(session).order_by(Expense.expense_date.desc()).all()
        finally:
            session.close()

    def get_expense_by_id(self, expense_id: int) -> Expense | None:
        session = get_session()
        try:
            return self._query(session).filter(Expense.id == expense_id).first()
        finally:
            session.close()

    def create_expense(
        self,
        project_id: int,
        expense_date: date,
        expense_type: str,
        amount,
        remarks: str = None,
        is_reimbursed: str = "未提交",
        travel_person: str = None,
        travel_start_date: date = None,
        travel_end_date: date = None,
        expense_category: str = None,
        assessment_fee=None,
        actual_cost=None,
    ) -> Expense:
        session = get_session()
        try:
            if not project_id or int(project_id) <= 0:
                raise ValidationError("项目不能为空")
            project = self._project(session, project_id)
            if not project:
                raise ValidationError("无权访问该项目")
            combined = assessment_fee is not None or actual_cost is not None
            if combined:
                assessment_fee = validate_positive_amount(assessment_fee or 0, "考核费用")
                actual_cost = validate_positive_amount(actual_cost or 0, "出差费用")
                if assessment_fee == 0 and actual_cost == 0:
                    raise ValidationError("考核费用和出差费用不能同时为 0")
                expense_category, expense_type = None, "费用"
                amount = assessment_fee + actual_cost
            else:
                expense_category, expense_type = self._normalise_category_type(
                    expense_category, expense_type
                )
                amount = validate_positive_amount(amount, "金额")
            travel_person = optional_text(travel_person, "出差人员", 100)
            validate_date_range(travel_start_date, travel_end_date, "出差开始", "出差结束")
            remarks = optional_text(remarks, "备注", EXPENSE_FIELD_LIMITS["remarks"])
            is_reimbursed = validate_choice(is_reimbursed, "报销状态", {"未提交", "已提交", "已报销"})

            expense = Expense(
                owner_id=project.owner_id or owner_id_for(self.current_user),
                project_id=project_id,
                expense_date=expense_date,
                travel_person=travel_person,
                travel_start_date=travel_start_date,
                travel_end_date=travel_end_date,
                expense_category=expense_category,
                expense_type=expense_type,
                amount=amount,
                assessment_fee=assessment_fee if combined else None,
                actual_cost=actual_cost if combined else None,
                remarks=remarks,
                is_reimbursed=is_reimbursed,
            )
            session.add(expense)
            session.commit()
            session.refresh(expense)
            return expense
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    def update_expense(self, expense_id: int, **kwargs) -> Expense | None:
        session = get_session()
        try:
            expense = self._query(session).filter(Expense.id == expense_id).first()
            if not expense:
                return None
            target_project_id = kwargs.get("project_id", expense.project_id)
            project = self._project(session, target_project_id)
            if not project:
                raise ValidationError("无权访问目标项目")
            kwargs = dict(kwargs)
            combined = "assessment_fee" in kwargs or "actual_cost" in kwargs
            if combined:
                assessment_fee = validate_positive_amount(kwargs.get("assessment_fee", 0) or 0, "考核费用")
                actual_cost = validate_positive_amount(kwargs.get("actual_cost", 0) or 0, "出差费用")
                if assessment_fee == 0 and actual_cost == 0:
                    raise ValidationError("考核费用和出差费用不能同时为 0")
                kwargs.update(
                    assessment_fee=assessment_fee,
                    actual_cost=actual_cost,
                    amount=assessment_fee + actual_cost,
                    expense_category=None,
                    expense_type="费用",
                )
            else:
                requested_category = kwargs.get("expense_category", expense.expense_category)
                requested_type = kwargs.get("expense_type", expense.expense_type)
                requested_category, requested_type = self._normalise_category_type(
                    requested_category, requested_type
                )
                kwargs["expense_category"] = requested_category
                kwargs["expense_type"] = requested_type
            for key, value in kwargs.items():
                if hasattr(expense, key):
                    if key == "expense_category" and value is not None:
                        value = validate_choice(value, "费用分类", self.EXPENSE_CATEGORIES)
                    elif key == "expense_type":
                        value = validate_choice(value, "费用类型", self.EXPENSE_TYPES)
                    elif key == "travel_person":
                        value = optional_text(value, "出差人员", 100)
                    elif key in {"travel_start_date", "travel_end_date"}:
                        value = value or None
                    elif key in {"amount", "assessment_fee", "actual_cost"}:
                        value = validate_positive_amount(value, "金额")
                    elif key == "remarks":
                        value = optional_text(value, "备注", EXPENSE_FIELD_LIMITS["remarks"])
                    elif key == "is_reimbursed":
                        value = validate_choice(value, "报销状态", {"未提交", "已提交", "已报销"})
                    setattr(expense, key, value)
            validate_date_range(
                expense.travel_start_date,
                expense.travel_end_date,
                "出差开始",
                "出差结束",
            )
            session.commit()
            session.refresh(expense)
            return expense
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    def delete_expense(self, expense_id: int) -> bool:
        session = get_session()
        try:
            expense = self._query(session).filter(Expense.id == expense_id).first()
            if not expense:
                return False
            session.delete(expense)
            session.commit()
            return True
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    def delete_all_expenses(self) -> int:
        session = get_session()
        try:
            count = self._query(session).count()
            self._query(session).delete(synchronize_session=False)
            session.commit()
            return count
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    def get_total_amount(self, project_id: int | None):
        session = get_session()
        try:
            query = self._query(session)
            if project_id:
                query = query.filter(Expense.project_id == project_id)
            return sum(self.amount_breakdown(expense)["actual_cost"] for expense in query.all())
        finally:
            session.close()

    def get_profit_summary(self, project_id: int | None, search: str = "", owner_id=None) -> dict:
        """Return assessment fees, actual travel costs, and realized profit."""
        session = get_session()
        try:
            query = self._filtered_expense_query(session, project_id, search, owner_id)
            rows = query.all()
            breakdowns = [(row, self.amount_breakdown(row)) for row in rows]
            assessment = sum(item["assessment_fee"] for _row, item in breakdowns)
            actual_cost = sum(item["actual_cost"] for _row, item in breakdowns)
            reimbursed = sum(
                item["actual_cost"]
                for row, item in breakdowns
                if row.is_reimbursed == "已报销"
            )
            return {
                "assessment_fee": assessment,
                "actual_cost": actual_cost,
                "profit": assessment - actual_cost,
                "reimbursed": reimbursed,
                "count": len(rows),
            }
        finally:
            session.close()

    def get_profit_summaries_by_project(self, project_id: int | None = None, search: str = "", owner_id=None) -> list[dict]:
        """Return one profit summary row per project."""
        session = get_session()
        try:
            rows = (
                self._filtered_expense_query(session, project_id, search, owner_id)
                .outerjoin(Project, Project.id == Expense.project_id)
                .with_entities(Expense, Project.name.label("project_name"))
                .order_by(Project.name, Expense.expense_date)
                .all()
            )
            grouped = {}
            for expense, project_name in rows:
                key = expense.project_id or 0
                if key not in grouped:
                    grouped[key] = {
                        "project_id": expense.project_id,
                        "owner_id": expense.owner_id,
                        "project_name": project_name or "未关联项目",
                        "assessment_fee": 0.0,
                        "actual_cost": 0.0,
                        "profit": 0.0,
                        "count": 0,
                    }
                item = grouped[key]
                breakdown = self.amount_breakdown(expense)
                item["assessment_fee"] += breakdown["assessment_fee"]
                item["actual_cost"] += breakdown["actual_cost"]
                item["profit"] = item["assessment_fee"] - item["actual_cost"]
                item["count"] += 1

            if project_id and not grouped:
                project = self._project(session, project_id)
                if project:
                    grouped[project.id] = {
                        "project_id": project.id,
                        "owner_id": project.owner_id,
                        "project_name": project.name,
                        "assessment_fee": 0.0,
                        "actual_cost": 0.0,
                        "profit": 0.0,
                        "count": 0,
                    }
            return list(grouped.values())
        finally:
            session.close()

    def export_expenses(self, project_id: int, file_path: str) -> str:
        expenses = self.get_expenses_by_project(project_id)
        if not expenses:
            raise ValidationError("当前项目暂无费用可导出")
        target = Path(file_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "出差费用"
        headers = [
            "费用日期", "出差人员", "出差开始", "出差结束",
            "费用分类", "考核费用", "出差费用", "本条利润", "备注", "报销状态",
        ]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DCE8F8")
        for expense in expenses:
            breakdown = self.amount_breakdown(expense)
            sheet.append([
                expense.expense_date,
                expense.travel_person or "",
                expense.travel_start_date,
                expense.travel_end_date,
                expense.expense_category or "合并费用",
                breakdown["assessment_fee"],
                breakdown["actual_cost"],
                breakdown["profit"],
                expense.remarks or "",
                expense.is_reimbursed or "",
            ])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(sheet.columns, start=1):
            width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[get_column_letter(index)].width = width
        workbook.save(target)
        return str(target)
