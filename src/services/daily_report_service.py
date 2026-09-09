from typing import List, Optional
from datetime import date
from src.models import DailyReport, get_session
from src.utils.validators import (
    DAILY_FIELD_LIMITS,
    ValidationError,
    optional_text,
    require_text,
    validate_choice,
)
from pathlib import Path

try:
    from docx import Document
except ImportError:  # pragma: no cover - dependency is part of requirements.txt
    Document = None

class DailyReportService:
    def get_reports_page(self, page=1, page_size=20, project_id=None):
        session = get_session()
        try:
            query = session.query(DailyReport)
            if project_id:
                query = query.filter(DailyReport.project_id == project_id)
            query = query.order_by(DailyReport.report_date.desc())
            return query.count(), query.offset((page - 1) * page_size).limit(page_size).all()
        finally:
            session.close()
    def get_reports_by_project(self, project_id: int) -> List[DailyReport]:
        session = get_session()
        try:
            return (
                session.query(DailyReport)
                .filter(DailyReport.project_id == project_id)
                .order_by(DailyReport.report_date.desc())
                .all()
            )
        finally:
            session.close()

    def get_report_by_id(self, report_id: int) -> DailyReport | None:
        session = get_session()
        try:
            return session.query(DailyReport).filter(DailyReport.id == report_id).first()
        finally:
            session.close()

    def create_report(
        self,
        project_id: int,
        report_date: date,
        work_content: str = None,
        problems: str = None,
        solutions: str = None,
        next_plan: str = None,
        remarks: str = None,
        status: str = "草稿",
    ) -> DailyReport:
        session = get_session()
        try:
            if not project_id or int(project_id) <= 0:
                raise ValidationError("项目不能为空")
            work_content = require_text(work_content, "SOP主题/步骤", DAILY_FIELD_LIMITS["work_content"])
            problems = optional_text(problems, "遇到的问题", DAILY_FIELD_LIMITS["problems"])
            solutions = optional_text(solutions, "解决方法", DAILY_FIELD_LIMITS["solutions"])
            next_plan = optional_text(next_plan, "后续建议", DAILY_FIELD_LIMITS["next_plan"])
            remarks = optional_text(remarks, "备注", DAILY_FIELD_LIMITS["remarks"])
            status = validate_choice(status, "SOP记录状态", {"草稿", "已完成"})

            existing = session.query(DailyReport).filter(
                DailyReport.project_id == project_id,
                DailyReport.report_date == report_date,
            ).first()
            if existing:
                existing.work_content = work_content
                existing.problems = problems
                existing.solutions = solutions
                existing.next_plan = next_plan
                existing.remarks = remarks
                existing.status = status
                session.commit()
                session.refresh(existing)
                return existing

            report = DailyReport(
                project_id=project_id,
                report_date=report_date,
                work_content=work_content,
                problems=problems,
                solutions=solutions,
                next_plan=next_plan,
                remarks=remarks,
                status=status,
            )
            session.add(report)
            session.commit()
            session.refresh(report)
            return report
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    def update_report(self, report_id: int, **kwargs) -> DailyReport | None:
        session = get_session()
        try:
            report = session.query(DailyReport).filter(DailyReport.id == report_id).first()
            if not report:
                return None
            target_project_id = int(kwargs.get("project_id", report.project_id))
            target_date = kwargs.get("report_date", report.report_date)
            if not target_project_id:
                raise ValidationError("项目不能为空")
            if not target_date:
                raise ValidationError("SOP记录日期不能为空")
            duplicate = session.query(DailyReport).filter(
                DailyReport.project_id == target_project_id,
                DailyReport.report_date == target_date,
                DailyReport.id != report_id,
            ).first()
            if duplicate:
                raise ValidationError("该项目在所选日期已存在SOP记录")
            for key, value in kwargs.items():
                if hasattr(report, key):
                    if key == "work_content":
                        value = require_text(value, "SOP主题/步骤", DAILY_FIELD_LIMITS["work_content"])
                    elif key in ["problems", "solutions", "next_plan"]:
                        value = optional_text(value, {"solutions": "解决方法", "next_plan": "后续建议"}.get(key, key), DAILY_FIELD_LIMITS[key])
                    elif key == "remarks":
                        value = optional_text(value, "备注", DAILY_FIELD_LIMITS["remarks"])
                    elif key == "status":
                        value = validate_choice(value, "SOP记录状态", {"草稿", "已完成"})
                    setattr(report, key, value)
            session.commit()
            session.refresh(report)
            return report
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    def delete_report(self, report_id: int) -> bool:
        session = get_session()
        try:
            report = session.query(DailyReport).filter(DailyReport.id == report_id).first()
            if not report:
                return False
            session.delete(report)
            session.commit()
            return True
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    def delete_all_reports(self) -> int:
        session = get_session()
        try:
            count = session.query(DailyReport).count()
            session.query(DailyReport).delete(synchronize_session=False)
            session.commit()
            return count
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    def get_report_count(self) -> int:
        session = get_session()
        try:
            return session.query(DailyReport).count()
        finally:
            session.close()

    def export_reports(self, project_id: int, file_path: str, file_format: str) -> str:
        reports = self.get_reports_by_project(project_id)
        if not reports:
            raise ValidationError("当前项目暂无SOP问题记录可导出")
        fmt = file_format.lower().lstrip(".")
        if fmt == "md":
            content = ["# SOP问题与解决方法", ""]
            for report in reversed(reports):
                content.extend([
                    f"## {report.report_date}",
                    f"- 状态：{report.status or ''}",
                    f"- SOP主题/步骤：{report.work_content or ''}",
                    f"- 遇到的问题：{report.problems or ''}",
                    f"- 解决方法：{report.solutions or ''}",
                    f"- 后续建议：{report.next_plan or ''}",
                    f"- 备注：{report.remarks or ''}",
                    "",
                ])
            Path(file_path).write_text("\n".join(content), encoding="utf-8")
            return file_path
        if fmt == "docx":
            if Document is None:
                raise ValidationError("未安装 python-docx，无法导出 Word")
            document = Document()
            document.add_heading("SOP问题与解决方法", level=1)
            for report in reversed(reports):
                document.add_heading(str(report.report_date), level=2)
                for label, value in [
                    ("状态", report.status), ("SOP主题/步骤", report.work_content),
                    ("遇到的问题", report.problems), ("解决方法", report.solutions),
                    ("后续建议", report.next_plan), ("备注", report.remarks),
                ]:
                    document.add_paragraph(f"{label}：{value or ''}")
            document.save(file_path)
            return file_path
        raise ValidationError("仅支持 Word (.docx) 或 Markdown (.md) 格式")
