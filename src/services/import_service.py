import json
from pathlib import Path
from openpyxl import load_workbook
from csv import DictReader
from src.utils.field_mapper import normalize_field_name
from src.utils.field_mapper import normalize_machine_header
from src.models import ImportFile, ImportRow, Project, get_session
from src.services.access_service import owner_id_for, scope_query
from src.services.machine_service import MachineService
from src.services.project_service import ProjectService
from src.utils.validators import MACHINE_FIELD_LIMITS, ValidationError, validate_ipv4
from src.ui.widgets.common import display_project_name
import re

IP_RE = re.compile(r"(?<!\d)(?:25[0-5]|2[0-4]\d|1?\d?\d)(?:\.(?:25[0-5]|2[0-4]\d|1?\d?\d)){3}(?!\d)")

class ImportService:
    def __init__(self, current_user=None):
        self.current_user = current_user
        self.last_sheet_scan = []

    def get_import_history(self, limit: int = 50, page: int = 1, page_size: int | None = None):
        session = get_session()
        try:
            query = (
                scope_query(session.query(ImportFile), ImportFile, self.current_user)
                .outerjoin(Project, Project.id == ImportFile.project_id)
                .with_entities(ImportFile, Project.name, Project.project_code)
                .order_by(ImportFile.imported_at.desc())
            )
            total = query.count()
            query = query.limit(page_size or limit)
            if page_size:
                query = query.offset((page - 1) * page_size)
            rows = query.all()
            result = [
                {
                    "project": display_project_name(project_name or "未关联项目"),
                    "file_name": record.file_name,
                    "records": record.records_count,
                    "new": record.new_count,
                    "merge": record.merge_count,
                    "conflict": record.conflict_count,
                    "status": record.status,
                    "imported_at": record.imported_at.strftime("%Y-%m-%d %H:%M:%S") if record.imported_at else "",
                }
                for record, project_name, project_code in rows
            ]
            return (total, result) if page_size else result
        finally:
            session.close()

    def read_file(self, file_path: str):
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            return self._read_xlsx(path)
        if suffix == ".csv":
            return self._read_csv(path)
        raise ValidationError("仅支持 Excel (.xlsx) 和 CSV (.csv) 文件")

    def read_file_groups(self, file_path: str) -> list[dict]:
        """Read one or more machine tables, preserving XLSX worksheet names."""
        path = Path(file_path)
        if path.suffix.lower() == ".csv":
            raw_rows = self._read_csv_source_rows(path)
            rows = [
                self._build_machine_row(
                    [normalize_machine_header(key) for key in raw_row.keys()],
                    list(raw_row.values()),
                    list(raw_row.keys()),
                )
                for raw_row in raw_rows
            ]
            rows = [row for row in rows if row]
            self.last_sheet_scan = [{"sheet_name": "CSV数据", "records": len(rows), "status": "已识别，待导入"}]
            return [{
                "sheet_name": "CSV数据",
                "rows": rows,
                "source_rows": raw_rows,
                "preamble_rows": [],
            }]
        if path.suffix.lower() != ".xlsx":
            raise ValidationError("仅支持 Excel (.xlsx) 和 CSV (.csv) 文件")
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            groups = []
            self.last_sheet_scan = []
            workbook_metadata = {}
            for worksheet in workbook.worksheets:
                rows = list(worksheet.iter_rows(values_only=True))
                if not rows or not any(value not in (None, "") for row in rows for value in row):
                    self.last_sheet_scan.append({
                        "sheet_name": worksheet.title, "records": 0, "status": "已跳过：空白工作表"
                    })
                    continue
                metadata = self._extract_project_metadata(rows)
                for key, value in metadata.items():
                    workbook_metadata.setdefault(key, value)
                try:
                    header_index, headers = self._find_header_row(rows)
                    source_headers = list(rows[header_index])
                except ValidationError as exc:
                    # Non-machine sheets are still retained losslessly.
                    source_rows = [
                        self._source_row([], raw)
                        for raw in rows
                        if any(value not in (None, "") for value in raw)
                    ]
                    groups.append({
                        "sheet_name": worksheet.title,
                        "rows": [],
                        "source_rows": source_rows,
                        "preamble_rows": [],
                        "project_metadata": {**workbook_metadata, **metadata},
                    })
                    self.last_sheet_scan.append({
                        "sheet_name": worksheet.title,
                        "records": len(source_rows),
                        "status": f"已读取 {len(source_rows)} 行，原始内容已保存",
                    })
                    continue
                result = []
                source_rows = [
                    raw for raw in rows[header_index + 1:]
                    if any(value not in (None, "") for value in raw)
                ]
                source_records = [self._source_row(source_headers, raw) for raw in source_rows]
                preamble_rows = [
                    {"行号": index + 1, **self._source_row([], raw)}
                    for index, raw in enumerate(rows[:header_index])
                    if any(value not in (None, "") for value in raw)
                ]
                if any(value not in (None, "") for value in source_headers):
                    preamble_rows.append({
                        "行号": header_index + 1,
                        "行类型": "原始表头",
                        **self._source_row([], source_headers),
                    })
                for raw in source_rows:
                    item = self._build_machine_row(headers, raw, source_headers)
                    if item:
                        result.append(item)
                if result or source_rows:
                    groups.append({
                        "sheet_name": worksheet.title,
                        "rows": result,
                        "source_rows": source_records,
                        "preamble_rows": preamble_rows,
                        "project_metadata": {**workbook_metadata, **metadata},
                    })
                    self.last_sheet_scan.append({
                        "sheet_name": worksheet.title,
                        "records": len(source_records),
                        "status": f"已读取 {len(source_rows)} 行，机器记录 {len(result)} 条，待导入",
                    })
                else:
                    self.last_sheet_scan.append({
                        "sheet_name": worksheet.title, "records": 0, "status": "已读取，无数据行"
                    })
            if not groups:
                raise ValidationError(
                    "未识别到任何包含机器信息的工作表，请确认至少一个子表包含角色及业务网IP/集群网IP等字段"
                )
            for group in groups:
                group["project_metadata"] = {**workbook_metadata, **group.get("project_metadata", {})}
            return groups
        finally:
            workbook.close()

    def get_project_label(self, project_id: int) -> str:
        session = get_session()
        try:
            project = scope_query(session.query(Project), Project, self.current_user).filter(
                Project.id == project_id
            ).first()
            return display_project_name(project.name) if project else "未知项目"
        finally:
            session.close()

    def get_or_create_import_project(self, sheet_name: str, file_stem: str, index: int, metadata: dict | None = None) -> int:
        """Match a project by the worksheet customer name, creating it when absent."""
        source_name = (sheet_name or file_stem or f"导入客户{index + 1}").strip()[:100]
        parts = re.split(r"\s*[-－—_/\\]\s*", source_name, maxsplit=1)
        project_name = parts[0].strip() or source_name
        sales_name = parts[1].strip() if len(parts) > 1 else ""
        # A sheet named "湘南学院-李华敏" is treated as project "湘南学院"
        # with the suffix mapped to the salesperson; do not duplicate it into 客户.
        customer_name = "" if sales_name else project_name
        metadata = metadata or {}
        sales_name = (metadata.get("sales") or sales_name).strip()[:100]
        location_name = str(metadata.get("location") or "").strip()[:100]
        remarks_text = str(metadata.get("remarks") or "").strip()[:4000]
        session = get_session()
        try:
            # The worksheet/project name is the primary association key. Matching
            # by customer first can incorrectly merge different projects owned by
            # the same customer and then make their IPs collide.
            project_query = scope_query(session.query(Project), Project, self.current_user)
            existing = project_query.filter(Project.name == project_name).first()
            if not existing and source_name != project_name:
                existing = project_query.filter(Project.name == source_name).first()
                if existing:
                    existing.name = project_name
                    existing.customer = customer_name
                    existing.sales = sales_name or existing.sales
                    existing.location = location_name or existing.location
                    existing.remarks = remarks_text or existing.remarks
                    session.commit()
            if not existing and customer_name:
                # Compatibility fallback for legacy projects whose name was
                # imported differently; only use an unambiguous customer match.
                customer_matches = project_query.filter(
                    Project.customer == customer_name
                ).limit(2).all()
                if len(customer_matches) == 1:
                    existing = customer_matches[0]
            if existing:
                if sales_name:
                    existing.customer = ""
                if sales_name:
                    existing.sales = sales_name
                if location_name and not existing.location:
                    existing.location = location_name
                if remarks_text and not existing.remarks:
                    existing.remarks = remarks_text
                if sales_name or location_name or remarks_text:
                    session.commit()
                return existing.id
        finally:
            session.close()
        project = ProjectService(self.current_user).create_project(
                name=project_name,
                customer=customer_name,
                sales=sales_name or None,
                    location=location_name or None,
                    project_code=None,
                    status="待开始",
                    remarks=remarks_text or None,
                )
        return project.id

    @staticmethod
    def _extract_project_metadata(rows):
        """Read project-level sales/location values from label-value cells."""
        aliases = {
            "sales": {"销售", "销售人员", "销售姓名", "sales"},
            "location": {"地点", "实施地点", "项目地点", "所在地", "location"},
            "remarks": {"备注", "说明", "remark", "remarks"},
        }
        label_values = set().union(*aliases.values()) | {"实施时间", "日期", "时间"}
        metadata = {}
        limited_rows = rows[:80]
        for row_index, row in enumerate(limited_rows):
            values = ["" if value is None else str(value).strip() for value in row]
            for index, value in enumerate(values):
                compact = re.sub(r"[\s:：()（）]", "", value).lower()
                for field, names in aliases.items():
                    if compact not in names or field in metadata:
                        continue
                    candidates = [
                        item for item in values[index + 1:] + values[:index]
                        if item and item.lower() not in label_values
                    ]
                    if not any(candidates):
                        candidates = [
                            "" if index >= len(next_row) or next_row[index] is None
                            else str(next_row[index]).strip()
                            for next_row in limited_rows[row_index + 1:row_index + 4]
                        ]
                    candidate = next((item for item in candidates if item), "")
                    if candidate and not IP_RE.fullmatch(candidate):
                        metadata[field] = candidate
        return metadata

    @staticmethod
    def combine_results(results: list[dict]) -> dict:
        combined = {"records": 0, "new": 0, "merge": 0, "conflict": 0, "errors": [], "conflicts": []}
        for result in results:
            for key in ("records", "new", "merge", "conflict"):
                combined[key] += result.get(key, 0)
            combined["errors"].extend(result.get("errors", []))
            combined["conflicts"].extend(result.get("conflicts", []))
        return combined

    def _read_xlsx(self, path: Path):
        wb = load_workbook(path, data_only=True, read_only=True)
        try:
            rows = list(wb.active.iter_rows(values_only=True))
            if not rows:
                return []
            header_index, headers = self._find_header_row(rows)
            result = []
            for raw in rows[header_index + 1:]:
                if self._header_score(raw) >= 2:
                    continue
                item = self._build_machine_row(headers, raw, rows[header_index])
                if item:
                    result.append(item)
            return result
        finally:
            wb.close()

    def _read_csv(self, path: Path):
        result = []
        for row in self._read_csv_source_rows(path):
            source_headers = list(row.keys())
            item = self._build_machine_row(
                [normalize_machine_header(key) for key in source_headers],
                list(row.values()),
                source_headers,
            )
            if item:
                result.append(item)
        return result

    @staticmethod
    def _read_csv_source_rows(path: Path):
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return [dict(row) for row in DictReader(f)]

    def _find_header_row(self, rows):
        network_fields = {"ip", "business_ip", "cluster_ip", "compute_ip", "storage_ip"}
        candidates = []
        for index, row in enumerate(rows[:200]):
            headers = [normalize_machine_header(value) for value in row]
            fields = {field for mapped in headers for field in mapped}
            candidates.append((bool(fields & network_fields), len(fields), index, headers))
        has_network, _score, header_index, headers = max(
            candidates, key=lambda item: (item[0], item[1]), default=(False, 0, 0, [])
        )
        if not has_network:
            raise ValidationError("未找到 IP/网络列")
        return header_index, headers

    @staticmethod
    def _header_score(row) -> int:
        known = 0
        for value in row:
            if normalize_machine_header(value):
                known += 1
        return known

    @staticmethod
    def _build_machine_row(headers, raw, source_headers=None):
        item = {}
        extra_info = {}
        for index, value in enumerate(raw):
            if index >= len(headers) or value in (None, ""):
                continue
            fields = headers[index]
            if not fields:
                label = str((source_headers or [])[index] or f"列{index + 1}") if index < len(source_headers or []) else f"列{index + 1}"
                extra_info[label] = value
                continue
            text_value = str(value).strip()
            ips = IP_RE.findall(text_value)
            for field in fields:
                if field in {"ip", "business_ip", "cluster_ip", "compute_ip", "storage_ip"}:
                    if ips:
                        item[field] = ips[0]
                elif field == "gpu_count" and text_value in {"无", "无GPU", "-"}:
                    continue
                elif field == "gpu_model" and text_value in {"无", "-"}:
                    continue
                else:
                    item[field] = value
        if extra_info:
            item["extra_info"] = json.dumps(extra_info, ensure_ascii=False, default=str)
        # Many implementation sheets do not have a management IP column.
        if not item.get("ip"):
            item["ip"] = next(
                (item[field] for field in ("business_ip", "cluster_ip", "compute_ip", "storage_ip") if item.get(field)),
                "",
            )
        if not item["ip"]:
            return {}
        return item

    @staticmethod
    def _source_row(headers, raw):
        result = {}
        for index, value in enumerate(raw):
            if value in (None, ""):
                continue
            label = headers[index] if index < len(headers) and headers[index] not in (None, "") else f"列{index + 1}"
            label = str(label)
            if label in result:
                label = f"{label}（列{index + 1}）"
            result[label] = value
        return result

    def import_machines(self, project_id: int, rows: list[dict], file_name: str,
                        file_type: str = "", file_size: int = 0,
                        source_rows: list[dict] | None = None,
                        sheet_name: str = "导入数据") -> dict:
        """将导入记录按项目/IP写入机器表，已有机器只补充空字段。"""
        if not project_id:
            raise ValidationError("请先选择项目")

        session = get_session()
        try:
            project = scope_query(session.query(Project), Project, self.current_user).filter(
                Project.id == project_id
            ).first()
            if not project:
                raise ValidationError("目标项目不存在，请刷新项目列表后重新选择")
            project_name = project.name
        finally:
            session.close()

        machine_service = MachineService(self.current_user)
        allowed_fields = set(MACHINE_FIELD_LIMITS) | {"account", "extra_info"}
        summary = {"records": len(rows), "new": 0, "merge": 0, "conflict": 0, "errors": []}
        conflicts = []

        for row_number, raw_row in enumerate(rows, start=2):
            row = {key: value for key, value in raw_row.items() if key in allowed_fields or key == "ip"}
            try:
                raw_ip = row.pop("ip", "")
                # A spreadsheet row without any network IP is still valid source
                # data. It is persisted in ImportRow below and must not be
                # reported as a failed import merely because it cannot become a
                # Machine row yet.
                if not raw_ip:
                    continue
                ip = validate_ipv4(raw_ip, "IP", required=True)
                machine, is_new, row_conflicts = machine_service.merge_machine_data(
                    project_id, ip, row, require_network=False
                )
                if is_new:
                    summary["new"] += 1
                else:
                    summary["merge"] += 1
                for field, values in row_conflicts.items():
                    summary["conflict"] += 1
                    conflicts.append({
                        "row": row_number,
                        "ip": ip,
                        "field": field,
                        "old": values["old"],
                        "new": values["new"],
                        "action": "已保留原值",
                    })
            except Exception as exc:
                summary["errors"].append({"row": row_number, "message": str(exc), "data": dict(raw_row), "project_id": project_id})

        session = get_session()
        try:
            record = ImportFile(
                owner_id=project.owner_id or owner_id_for(self.current_user),
                project_id=project_id,
                file_name=file_name,
                file_type=file_type,
                file_size=file_size,
                records_count=summary["records"],
                new_count=summary["new"],
                merge_count=summary["merge"],
                conflict_count=summary["conflict"],
                status="已完成" if not summary["errors"] else "部分完成",
            )
            session.add(record)
            session.flush()
            for row_number, raw_row in enumerate(source_rows or rows, start=2):
                machine_row = self._build_machine_row(
                    [normalize_machine_header(key) for key in raw_row.keys()],
                    list(raw_row.values()),
                    list(raw_row.keys()),
                ) if isinstance(raw_row, dict) else None
                session.add(ImportRow(
                    owner_id=project.owner_id or owner_id_for(self.current_user),
                    project_id=project_id,
                    file_name=file_name,
                    sheet_name=sheet_name,
                    row_number=row_number,
                    row_json=json.dumps(raw_row, ensure_ascii=False, default=str),
                    import_status="机器已写入" if machine_row and machine_row.get("ip") else "原始数据已保存",
                ))
            session.commit()
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

        return {**summary, "conflicts": conflicts, "project_id": project_id, "project_name": project_name}
